"""Live integration verification for Multimodal Ingestion and Document Processing."""

import io
import httpx
try:
    from pptx import Presentation  # type: ignore
except ImportError:
    Presentation = None

BASE_URL = "http://127.0.0.1:8000"


def run_live_test():
    with httpx.Client(base_url=BASE_URL, timeout=15.0) as client:
        print("[1] Authenticating as demo student...")
        login_res = client.post(
            "/api/v1/auth/login",
            json={"email": "sarah@knovara.edu", "password": "StrongPassword2026!"},
        )
        assert login_res.status_code == 200, f"Login failed: {login_res.text}"
        token = login_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}
        print("  -> Authenticated successfully.")

        print("[2] Fetching course workspaces...")
        courses_res = client.get("/api/v1/courses", headers=headers)
        assert courses_res.status_code == 200, f"Get courses failed: {courses_res.text}"
        courses = courses_res.json()
        assert len(courses) > 0, "No courses found for student."
        course_id = courses[0]["id"]
        course_name = courses[0]["name"]
        print(f"  -> Using course '{course_name}' ({course_id})")

        print("[3] Seeding canonical multimodal demo materials...")
        seed_res = client.post(
            f"/api/v1/courses/{course_id}/documents/demo-seed", headers=headers
        )
        assert seed_res.status_code == 201, f"Seed materials failed: {seed_res.text}"
        seeded_docs = seed_res.json()
        print(f"  -> Successfully seeded {len(seeded_docs)} multimodal documents.")

        print("[4] Inspecting PDF textbook and citations...")
        pdf_meta = next(d for d in seeded_docs if d["filename"] == "ML_Textbook_Chapter4_DecisionTrees.pdf")
        pdf_detail = client.get(
            f"/api/v1/courses/{course_id}/documents/{pdf_meta['id']}", headers=headers
        ).json()
        assert pdf_detail["processing_status"] == "COMPLETED"
        assert len(pdf_detail["chunks"]) == 3
        # Check canonical page 42 citation
        assert pdf_detail["chunks"][0]["page_number"] == 42
        print(f"  -> PDF Textbook verified: Page {pdf_detail['chunks'][0]['page_number']} citation aligned.")

        print("[5] Inspecting PPTX slides and citations...")
        pptx_meta = next(d for d in seeded_docs if d["filename"] == "ML_Lecture05_Ensembles_RandomForest.pptx")
        pptx_detail = client.get(
            f"/api/v1/courses/{course_id}/documents/{pptx_meta['id']}", headers=headers
        ).json()
        assert pptx_detail["chunks"][0]["slide_number"] == 18
        print(f"  -> PPTX Slides verified: Slide {pptx_detail['chunks'][0]['slide_number']} citation aligned.")

        print("[6] Inspecting Video transcript and timestamp citations...")
        video_meta = next(d for d in seeded_docs if d["filename"] == "ML_Lecture05_Entropy_Video_Transcript.mp4")
        video_detail = client.get(
            f"/api/v1/courses/{course_id}/documents/{video_meta['id']}", headers=headers
        ).json()
        assert video_detail["chunks"][0]["timestamp_start"] == "18:20"
        assert video_detail["chunks"][0]["timestamp_end"] == "20:05"
        print(f"  -> Video transcript verified: {video_detail['chunks'][0]['timestamp_start']} to {video_detail['chunks'][0]['timestamp_end']} citation aligned.")

        print("[7] Uploading live PPTX slide deck...")
        assert Presentation is not None, "python-pptx must be installed"
        prs = Presentation()
        s1 = prs.slides.add_slide(prs.slide_layouts[0])
        if s1.shapes.title:
            s1.shapes.title.text = "Decision Trees & Supervised Classification"
        s2 = prs.slides.add_slide(prs.slide_layouts[1])
        if s2.shapes.title:
            s2.shapes.title.text = "Entropy and Information Gain Principles"

        pptx_bytes = io.BytesIO()
        prs.save(pptx_bytes)
        pptx_bytes.seek(0)

        upload_res = client.post(
            f"/api/v1/courses/{course_id}/documents",
            files={
                "file": (
                    "Live_Lecture_Slides.pptx",
                    pptx_bytes,
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
            },
            headers=headers,
        )
        assert upload_res.status_code == 201, f"Upload failed: {upload_res.text}"
        uploaded_doc = upload_res.json()
        assert uploaded_doc["processing_status"] == "COMPLETED"
        assert len(uploaded_doc["chunks"]) >= 2
        print(f"  -> Uploaded live PPTX '{uploaded_doc['filename']}': extracted {len(uploaded_doc['chunks'])} chunks with slide citations.")

        print("\nAll live integration checks passed successfully! Multimodal ingestion pipeline operational.")


if __name__ == "__main__":
    run_live_test()
